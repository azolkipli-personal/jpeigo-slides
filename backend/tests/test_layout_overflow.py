#!/usr/bin/env python3
"""Layout-overflow checks: length-aware scaling for every target language, plus
the neutralising of the PDF-import autofit/wrap artifacts on injected frames.

Standalone — no pytest required:

    cd backend && venv/bin/python tests/test_layout_overflow.py

Why these four: the measured causes of the JP->EN deck overflow were (1) a font
scale that only ever ran for EN->JA, (2) every frame left on spAutoFit so the box
grew over its neighbour, (3) wrap left off so the overflow ran off the slide
edge, and (4) frames the translation did not disturb being fair game for a
layout change nobody asked for.
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Pt

from app.core.injector import (
    calculate_font_scale,
    frame_text_exceeds_box,
    frame_text_grew,
    paragraph_font_scale,
    replace_text_in_shape,
    resolve_font_size,
)
from app.models import TranslatedRun

FAILURES = []


def check(name, ok, detail=''):
    print(('  PASS  ' if ok else '  FAIL  ') + name + (f'  — {detail}' if detail else ''))
    if not ok:
        FAILURES.append(name)


def run_of(run_id, original, translated, target='en'):
    return TranslatedRun(
        run_id=run_id,
        original_text=original,
        translated_text=translated,
        source_language='ja' if target == 'en' else 'en',
        target_language=target,
        model_used='test',
    )


def pdf_import_frame(source_text, font_pt=12.0, width_pt=400.0, height_pt=60.0):
    """A text frame shaped the way LibreOffice's PDF import writes them:
    wrap off and spAutoFit on."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Pt(0), Pt(0), Pt(width_pt), Pt(height_pt))
    frame = box.text_frame
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    run = frame.paragraphs[0].add_run()
    run.text = source_text
    run.font.size = Pt(font_pt)
    return box, frame


def inject(box, runs):
    failed = replace_text_in_shape(box, runs, 0, '0')
    assert failed == [], f'unexpected injection failures: {failed}'


print('\n[1] a frame whose text grows gets a non-growing autofit')

box, frame = pdf_import_frame('Hello world', width_pt=400.0)
check('pre-fix: the PDF-import frame starts on spAutoFit with wrap off',
      frame.auto_size == MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT and frame.word_wrap is False,
      f'auto_size={frame.auto_size} word_wrap={frame.word_wrap}')

runs = [run_of('run_0_0_0_0_1', 'Hello world', 'Hello there my dear friend')]
inject(box, runs)
check('the frame is classified as grown (and as still fitting its box)',
      frame_text_grew(runs) and not frame_text_exceeds_box(frame, box),
      f'grew={frame_text_grew(runs)} exceeds={frame_text_exceeds_box(frame, box)}')

check('post-fix: the grown frame is off spAutoFit',
      frame.auto_size != MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT, f'auto_size={frame.auto_size}')
check('post-fix: the replacement is normAutofit (shrink text, never grow the box)',
      frame.auto_size == MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE, f'auto_size={frame.auto_size}')
check('the translation really did grow (and was written)',
      frame.paragraphs[0].runs[0].text == 'Hello there my dear friend',
      repr(frame.paragraphs[0].runs[0].text))
check('wrap is untouched while the text still fits the box width',
      frame.word_wrap is False, f'word_wrap={frame.word_wrap}')


print('\n[2] a frame that needs wrap gets word_wrap=True')

box, frame = pdf_import_frame('短い', font_pt=18.0, width_pt=120.0)
long_translation = ('This English translation is far longer than the original '
                    'Japanese label ever was')
inject(box, [run_of('run_0_0_0_0_1', '短い', long_translation)])
check('the frame is classified as needing wrap',
      frame_text_exceeds_box(frame, box),
      f'{len(long_translation)} chars at 18pt in a 120pt box')
check('post-fix: word_wrap is turned on', frame.word_wrap is True,
      f'word_wrap={frame.word_wrap}')
check('post-fix: a frame that no longer fits is also off spAutoFit',
      frame.auto_size == MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE, f'auto_size={frame.auto_size}')

box, frame = pdf_import_frame('Source text that already fits its box', width_pt=600.0)
inject(box, [run_of('run_0_0_0_0_1', 'Source text that already fits its box',
                    'Fits')])
check('a translation that still fits is not given wrap',
      frame.word_wrap is False, f'word_wrap={frame.word_wrap}')


print('\n[3] a frame unchanged by translation is left alone')

box, frame = pdf_import_frame('Governance framework', width_pt=400.0)
inject(box, [run_of('run_0_0_0_0_1', 'Governance framework', 'Governance')])

check('the shorter translation was written',
      frame.paragraphs[0].runs[0].text == 'Governance',
      repr(frame.paragraphs[0].runs[0].text))
check('post-fix: autofit is left exactly as the source had it',
      frame.auto_size == MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT, f'auto_size={frame.auto_size}')
check('post-fix: wrap is left exactly as the source had it',
      frame.word_wrap is False, f'word_wrap={frame.word_wrap}')
check('no size stamp is written either',
      frame.paragraphs[0].runs[0].font.size == Pt(12.0),
      f'size={frame.paragraphs[0].runs[0].font.size.pt if frame.paragraphs[0].runs[0].font.size else None}')


print('\n[4] EN->JA scaling output is unchanged for a known input')

# Pinned to the values the tuned parameters produce (cjk_ratio 2.2, 5% tolerance,
# 0.5 floor): any edit to those numbers fails here. 'Workshop ' is 9.0 units wide
# and 'ワークショップ' (7 CJK) 15.4, so the ratio is 9/15.4.
known = paragraph_font_scale([run_of('run_3_4_2_0_17', 'Workshop ', 'ワークショップ',
                                     target='ja')])
check('EN->JA pair returns the exact original:translated ratio',
      abs(known - (9.0 / 15.4)) < 1e-9, f'scale={known}')
check('the EN->JA value equals calculate_font_scale on the same pair',
      abs(calculate_font_scale('Workshop ', 'ワークショップ') - (9.0 / 15.4)) < 1e-9)
floor = paragraph_font_scale([run_of('run_0_0_0_0_1', 'x', 'あ' * 10, target='ja')])
check('an EN->JA pair below the floor still stops at 0.5', floor == 0.5,
      f'scale={floor}')

# ...and the same computation now runs for a growing JP->EN pair — 'AI時代の経営'
# is 13.0 units wide, 'Management in the AI era' 24.0 — where the old gate
# returned an unconditional 1.0.
jp_en = paragraph_font_scale([run_of('run_0_0_0_0_1', 'AI時代の経営',
                                     'Management in the AI era', target='en')])
check('the scale gate no longer blocks non-ja targets',
      abs(jp_en - (13.0 / 24.0)) < 1e-9, f'scale={jp_en}')
check('a JP->EN translation that did not grow is still left at 1.0',
      paragraph_font_scale([run_of('run_3_9_2_0_40', 'Workshop ', 'Workshop',
                                   target='en')]) == 1.0)

# The scale has to reach the written font size, not just the number.
size = resolve_font_size(run_of('run_0_0_0_0_1', 'AI時代の経営',
                                'Management in the AI era', target='en'),
                         jp_en, Pt(20))
check('a growing JP->EN paragraph actually gets a smaller font size',
      size is not None and size < 20.0, f'{size}pt < 20.0pt')

print()
if FAILURES:
    print(f'{len(FAILURES)} CHECK(S) FAILED: ' + ', '.join(FAILURES))
    sys.exit(1)
print('All layout-overflow checks passed.')


def test_layout_overflow() -> None:
    """The module-level checks above run at import; expose their verdict to pytest."""
    assert not FAILURES, f'layout-overflow checks failed: {FAILURES}'
