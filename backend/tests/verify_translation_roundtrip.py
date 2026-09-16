#!/usr/bin/env python3
"""Deterministic format-integrity verification for translated PPTX output.

Two modes:

  Live round trip against a running backend:
      python tests/verify_translation_roundtrip.py --source test_real.pptx

  Offline comparison of any two decks:
      python tests/verify_translation_roundtrip.py --original a.pptx --translated b.pptx

Checks are structural and deterministic — no AI, no network in offline mode:

  * slide count preserved
  * every paragraph keeps its run count (runs neither merged nor dropped)
  * no run that held text was silently emptied
  * Japanese runs carry <a:ea> (without it the CJK font override is cosmetic)
  * runs within one paragraph share a font size (per-paragraph scale)
  * translation coverage: runs still holding source text are listed

Exit code is 0 when no hard violation is found, 1 otherwise.
"""
from __future__ import annotations

import argparse
import json
import sys
import uuid
import urllib.request
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.shapes.group import GroupShape
from pptx.shapes.graphfrm import GraphicFrame
from lxml import etree

A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'
DGM_NS = 'http://schemas.openxmlformats.org/drawingml/2006/diagram'
R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def _fmt_sizes(sizes) -> str:
    """Format a set of font sizes that may contain None (sorted() would raise on a mix)."""
    known = sorted(v for v in sizes if v is not None)
    return str(known) + (' +no-size' if None in sizes else '')


def has_cjk(text: str) -> bool:
    return any('\u3040' <= c <= '\u30ff' or '\u4e00' <= c <= '\u9fff' for c in text)


@dataclass
class RunRecord:
    text: str
    size: float | None
    has_ea: bool


@dataclass
class Deck:
    """Paragraphs keyed by (slide, shape path, paragraph index)."""

    paragraphs: dict[tuple[int, str, int], list[RunRecord]] = field(default_factory=dict)
    slides: int = 0

    def add(self, key, text, size, has_ea):
        self.paragraphs.setdefault(key, []).append(RunRecord(text, size, has_ea))


def _run_size(run) -> float | None:
    try:
        return run.font.size.pt if run.font.size else None
    except Exception:
        return None


def _run_has_ea(run) -> bool:
    try:
        return run._r.find(f'.//{{{A_NS}}}ea') is not None
    except Exception:
        return False


def _smartart_texts(shape) -> list[str]:
    """SmartArt text, in the same document order the extractor/injector use."""
    try:
        graphic_data = shape._element.find(f'.//{{{A_NS}}}graphicData')
        if graphic_data is None or DGM_NS not in graphic_data.get('uri', ''):
            return []
        rel_ids = graphic_data.find(f'.//{{{DGM_NS}}}relIds')
        if rel_ids is None:
            return []
        rel_id = rel_ids.get(f'{{{R_NS}}}dm')
        if not rel_id:
            return []
        part = shape.part.related_part(rel_id)
        xml = etree.fromstring(part.blob)
        return [(t.text or '') for t in xml.iter(f'{{{A_NS}}}t') if (t.text or '').strip()]
    except Exception:
        return []


def _walk_text_frame(text_frame, deck: Deck, slide_idx: int, key_prefix: str) -> None:
    for para_idx, para in enumerate(text_frame.paragraphs):
        runs = [r for r in para.runs if (r.text or '').strip()]
        if not runs:
            continue
        key = (slide_idx, key_prefix, para_idx)
        for run in runs:
            deck.add(key, run.text, _run_size(run), _run_has_ea(run))


def _walk(shapes, deck: Deck, slide_idx: int, prefix: str = '') -> None:
    for idx, shape in enumerate(shapes):
        path = f'{prefix}{idx}'
        if isinstance(shape, GroupShape):
            _walk(shape.shapes, deck, slide_idx, prefix=f'{path}.')
            continue
        if isinstance(shape, GraphicFrame):
            if getattr(shape, 'has_table', False) and shape.has_table:
                for row_idx, row in enumerate(list(shape.table.rows)):
                    for col_idx, cell in enumerate(row.cells):
                        _walk_text_frame(
                            cell.text_frame, deck, slide_idx,
                            f'{path}.table.{row_idx}.{col_idx}',
                        )
                continue
            texts = _smartart_texts(shape)
            if texts:
                key = (slide_idx, f'smartart.{path}', 0)
                for text in texts:
                    deck.add(key, text, None, True)
            continue
        if getattr(shape, 'has_text_frame', False) and shape.text_frame is not None:
            _walk_text_frame(shape.text_frame, deck, slide_idx, path)


def read_deck(path: str | Path) -> Deck:
    prs = Presentation(str(path))
    deck = Deck(slides=len(prs.slides))
    for slide_idx, slide in enumerate(prs.slides):
        _walk(slide.shapes, deck, slide_idx)
    return deck


def inspect_deck(deck: Deck) -> dict:
    """Layer 1 signals measurable from one deck with no original to compare against."""
    ja_runs = 0
    ja_without_ea = 0
    examples = []
    mixed_paragraphs = 0
    multi_run_paragraphs = 0

    for key, runs in deck.paragraphs.items():
        if len(runs) > 1:
            multi_run_paragraphs += 1
            sizes = {r.size for r in runs}
            if len(sizes) > 1 and None not in sizes:
                mixed_paragraphs += 1
        for record in runs:
            if has_cjk(record.text):
                ja_runs += 1
                if not record.has_ea:
                    ja_without_ea += 1
                    if len(examples) < 5:
                        examples.append(f'slide {key[0]} {key[1]} para {key[2]}: {record.text[:28]!r}')

    return {
        'slides': deck.slides,
        'paragraphs': len(deck.paragraphs),
        'runs': sum(len(v) for v in deck.paragraphs.values()),
        'ja_runs': ja_runs,
        'ja_without_ea': ja_without_ea,
        'multi_run_paragraphs': multi_run_paragraphs,
        'mixed_size_paragraphs': mixed_paragraphs,
        'examples': examples,
    }


def compare(original: Deck, translated: Deck) -> tuple[list[str], list[str], dict]:
    """Return (violations, warnings, stats)."""
    violations: list[str] = []
    warnings: list[str] = []

    if original.slides != translated.slides:
        violations.append(f'slide count changed: {original.slides} -> {translated.slides}')

    missing = set(original.paragraphs) - set(translated.paragraphs)
    if missing:
        violations.append(f'{len(missing)} paragraphs disappeared, e.g. {sorted(missing)[:3]}')

    emptied = 0
    leftovers = 0
    total_runs = 0
    ja_runs = 0
    ja_without_ea = 0
    size_splits: list[str] = []

    for key, orig_runs in original.paragraphs.items():
        new_runs = translated.paragraphs.get(key)
        if new_runs is None:
            continue
        total_runs += len(orig_runs)

        if len(new_runs) != len(orig_runs):
            violations.append(
                f'run count changed in slide {key[0]} shape {key[1]} para {key[2]}: '
                f'{len(orig_runs)} -> {len(new_runs)}'
            )
            continue

        for o, n in zip(orig_runs, new_runs):
            if o.text.strip() and not n.text.strip():
                emptied += 1
            if n.text.strip() and n.text == o.text and has_cjk(o.text) is False:
                leftovers += 1
            if has_cjk(n.text):
                ja_runs += 1
                if not n.has_ea:
                    ja_without_ea += 1

        # Runs of one paragraph should share a single font size when they started
        # out uniform — a per-run scale splits them.
        orig_sizes = {r.size for r in orig_runs}
        new_sizes = {r.size for r in new_runs}
        if len(orig_sizes) == 1 and len(new_sizes) > 1:
            size_splits.append(
                f'slide {key[0]} shape {key[1]} para {key[2]}: '
                f'one original size {_fmt_sizes(orig_sizes)} -> {_fmt_sizes(new_sizes)}'
            )

    if emptied:
        violations.append(f'{emptied} runs lost their text (were non-empty, now empty)')
    if size_splits:
        violations.append(
            f'{len(size_splits)} paragraphs gained mixed font sizes: ' + '; '.join(size_splits[:3])
        )
    if ja_without_ea:
        warnings.append(
            f'{ja_without_ea} of {ja_runs} Japanese runs have no <a:ea> '
            f'(CJK typeface override will not apply)'
        )
    if leftovers:
        warnings.append(f'{leftovers} runs still hold source text (untranslated)')

    stats = {
        'slides': translated.slides,
        'paragraphs': len(original.paragraphs),
        'runs': total_runs,
        'ja_runs': ja_runs,
        'ja_without_ea': ja_without_ea,
        'untranslated_runs': leftovers,
        'mixed_size_paragraphs': len(size_splits),
    }
    return violations, warnings, stats


# ---------------------------------------------------------------- live mode

def _post_multipart(url: str, filename: str, payload: bytes, timeout: int = 900) -> dict:
    boundary = f'----jpeigo{uuid.uuid4().hex}'
    body = b''.join([
        f'--{boundary}\r\n'.encode(),
        f'Content-Disposition: form-data; name="file"; filename="{filename}"\r\n'.encode(),
        b'Content-Type: application/vnd.openxmlformats-officedocument.presentationml.presentation\r\n\r\n',
        payload,
        f'\r\n--{boundary}--\r\n'.encode(),
    ])
    request = urllib.request.Request(
        url, data=body,
        headers={'Content-Type': f'multipart/form-data; boundary={boundary}'},
    )
    with urllib.request.urlopen(request, timeout=timeout) as response:
        return json.loads(response.read())


def _post_json(url: str, payload: dict, timeout: int = 1800):
    request = urllib.request.Request(
        url, data=json.dumps(payload).encode(),
        headers={'Content-Type': 'application/json'},
    )
    with urllib.request.urlopen(request, timeout=timeout) as response:
        return response.headers, response.read()


def live_roundtrip(base_url: str, source: Path, out_path: Path, model: str) -> Path:
    print(f'[1/3] uploading {source.name}')
    document = _post_multipart(f'{base_url}/api/upload', source.name, source.read_bytes())
    job_id = document['job_id']
    slides = document.get('slides', [])
    if not slides:
        raise SystemExit(f'upload returned no slides; top-level keys={list(document)}')
    runs = [r for slide in slides for box in slide.get('text_boxes', []) for r in box.get('runs', [])]
    if not runs:
        raise SystemExit(f'no runs extracted; slide 0 keys={list(slides[0])}')
    print(f'      job_id={job_id}  runs={len(runs)}  slides={len(slides)}')

    print(f'[2/3] translating {len(runs)} runs to ja via {model}')
    _, body = _post_json(f'{base_url}/api/translate', {
        'runs': runs,
        'source_language': 'en',
        'target_language': 'ja',
        'model': model,
        'job_id': job_id,
    })
    translated = json.loads(body).get('translated_runs', [])
    (out_path.parent / 'translated_runs.json').write_text(json.dumps(translated, ensure_ascii=False, indent=2))
    changed = sum(1 for r in translated if r['original_text'] != r['translated_text'])
    failed = [r for r in translated if not r.get('success', True)]
    unchanged = [r for r in translated if r['original_text'] == r['translated_text'] and r['original_text'].strip()]
    print(f'      translated_runs={len(translated)}  changed={changed}  '
          f'provider_failures={len(failed)}  identity={len(unchanged)}')
    if failed:
        for run in failed[:10]:
            print(f'        FAILED on all providers: {run["run_id"]}  {run["original_text"][:40]!r}')
        raise SystemExit(f'{len(failed)} run(s) failed on every provider — source text passed through')
    for run in unchanged[:25]:
        print(f'        left as-is: {run["run_id"]}  {run["original_text"][:40]!r}')

    print('[3/3] exporting')
    headers, pptx = _post_json(f'{base_url}/api/export', {
        'job_id': job_id,
        'filename': out_path.name,
    })
    out_path.write_bytes(pptx)
    print(
        f'      wrote {out_path} ({len(pptx)} bytes)  '
        f'X-Injection-Failed={headers.get("X-Injection-Failed")}  '
        f'X-Injection-Total={headers.get("X-Injection-Total")}'
    )
    return out_path


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--original', type=Path, help='source deck (offline mode)')
    parser.add_argument('--translated', type=Path, help='translated deck (offline mode)')
    parser.add_argument('--source', type=Path, help='source deck to run through the live backend')
    parser.add_argument('--inspect', type=Path, nargs='+',
                        help='report Layer 1 signals for one or more decks (no original needed)')
    parser.add_argument('--out', type=Path, default=Path('/tmp/jpeigo-verify/translated.pptx'))
    parser.add_argument('--base-url', default='http://127.0.0.1:8002')
    parser.add_argument('--model', default='gemini')
    args = parser.parse_args()

    if args.inspect:
        reports = [(path, inspect_deck(read_deck(path))) for path in args.inspect]
        keys = ['slides', 'paragraphs', 'runs', 'ja_runs', 'ja_without_ea',
                'multi_run_paragraphs', 'mixed_size_paragraphs']
        print('metric'.ljust(26) + ''.join(path.name[:26].rjust(28) for path, _ in reports))
        for key in keys:
            print(key.ljust(26) + ''.join(str(report[key]).rjust(28) for _, report in reports))
        for path, report in reports:
            if report['examples']:
                print()
                print(f'{path.name} — Japanese runs missing <a:ea> (first {len(report["examples"])}):')
                for line in report['examples']:
                    print(f'  {line}')
        return 0

    original_path = args.original
    translated_path = args.translated

    if args.source:
        translated_path = live_roundtrip(args.base_url, args.source, args.out, args.model)
        original_path = args.source
    elif not (original_path and translated_path):
        parser.error('give --source, or both --original and --translated')

    original = read_deck(original_path)
    translated = read_deck(translated_path)
    violations, warnings, stats = compare(original, translated)

    print()
    print(f'original   : {original_path}')
    print(f'translated : {translated_path}')
    print()
    for label, value in stats.items():
        print(f'  {label:26} {value}')
    print()

    if violations:
        print(f'VIOLATIONS ({len(violations)}):')
        for item in violations:
            print(f'  ✗ {item}')
    else:
        print('VIOLATIONS: none')

    if warnings:
        print(f'WARNINGS ({len(warnings)}):')
        for item in warnings:
            print(f'  ! {item}')

    return 1 if violations else 0


if __name__ == '__main__':
    sys.exit(main())
