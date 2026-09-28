#!/usr/bin/env python3
"""Unit checks for the PDF bridge (app.core.pdf_bridge).

Standalone — no pytest, no network, no API keys:

    cd backend && venv/bin/python tests/test_pdf_bridge.py

Why this exists: every PDF upload/export goes through one LibreOffice call, and
the two failure modes that matter are silent — a conversion that hangs forever
(wedged LibreOffice) and concurrent conversions sharing one profile (the app
renders decks concurrently: preview route, QA, parallel uploads). Each check
below pins one contract of the wrapper:

  * the private profile dir is passed on every invocation and removed after it
  * timeout and non-zero exit surface as PdfBridgeError with a readable message
    and an HTTP status, not as a bare traceback
  * a missing soffice says so instead of raising FileNotFoundError
  * concurrent conversions use distinct profiles and all succeed

The happy-path conversions run the real LibreOffice (still no network), so this
also proves the bridge works on this machine end to end.
"""
import shutil
import subprocess
import sys
import tempfile
import threading
import types
from pathlib import Path

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

import app.core.pdf_bridge as pdf_bridge
from app.core.pdf_bridge import PdfBridgeError, pdf_to_pptx, pptx_to_pdf

SAMPLE = BACKEND.parent / 'public' / 'sample-deck.pptx'
PROFILE_PREFIX = '-env:UserInstallation=file://'
FAILURES: list[str] = []
CHECKS = 0


def check(name, ok, detail=''):
    global CHECKS
    CHECKS += 1
    print(('  PASS  ' if ok else '  FAIL  ') + name + (f'  — {detail}' if detail else ''))
    if not ok:
        FAILURES.append(name)


class SpySubprocess:
    """Stands in for the subprocess module: records every call, then either
    answers as configured (raise/returncode) or runs the real command."""

    def __init__(self):
        self.calls = []
        self.lock = threading.Lock()
        self.raise_timeout = False
        self.returncode = None
        self.stderr = b''

    def run(self, cmd, **kwargs):
        with self.lock:
            self.calls.append((list(cmd), dict(kwargs)))
        if self.raise_timeout:
            raise subprocess.TimeoutExpired(cmd, kwargs.get('timeout', 0))
        if self.returncode is not None:
            return types.SimpleNamespace(returncode=self.returncode,
                                         stdout=b'', stderr=self.stderr)
        return subprocess.run(cmd, **kwargs)


def patch_subprocess(spy):
    """Swap pdf_bridge's subprocess for the spy; returns the restore callable."""
    original = pdf_bridge.subprocess
    pdf_bridge.subprocess = types.SimpleNamespace(run=spy.run,
                                                  TimeoutExpired=subprocess.TimeoutExpired)
    return lambda: setattr(pdf_bridge, 'subprocess', original)


def profile_path(cmd) -> Path | None:
    for arg in cmd:
        if isinstance(arg, str) and arg.startswith(PROFILE_PREFIX):
            return Path(arg[len(PROFILE_PREFIX):])
    return None


def expect_bridge_error(name, fn, needle, status=None):
    """Run fn, expecting a PdfBridgeError whose message mentions `needle`."""
    try:
        fn()
    except PdfBridgeError as exc:
        ok = needle.lower() in str(exc).lower()
        status_ok = status is None or exc.status_code == status
        check(name, ok and status_ok,
              f'{exc} (status={exc.status_code}, wanted {status})')
    except Exception as exc:  # noqa: BLE001 — any other type is a contract break
        check(name, False, f'wrong exception: {type(exc).__name__}: {exc}')


def main() -> int:
    work = Path(tempfile.mkdtemp(prefix='pdfbridge-test-'))
    try:
        soffice = pdf_bridge.soffice_path()
        check('soffice resolves to an absolute path', bool(soffice), str(soffice))

        print('\n[1] happy path: pptx -> pdf -> pptx (real LibreOffice)')
        if not soffice or not SAMPLE.exists():
            print('  (skipping conversions: soffice or sample deck missing)')
        else:
            pdf = pptx_to_pdf(SAMPLE, work)
            check('pptx_to_pdf produced a PDF',
                  pdf.exists() and pdf.read_bytes()[:4] == b'%PDF', str(pdf.name))
            pptx = pdf_to_pptx(pdf, work)
            check('pdf_to_pptx names the intermediate after the whole source name',
                  pptx.name == f'{pdf.name}.pptx', str(pptx.name))
            from pptx import Presentation
            deck = Presentation(str(pptx))
            check('intermediate opens as a slide deck', len(deck.slides) > 0,
                  f'{len(deck.slides)} slides')

            print('\n[2] private profile: passed on every call, cleaned up after')
            spy = SpySubprocess()
            restore = patch_subprocess(spy)
            try:
                pdf_to_pptx(pdf, work)
            finally:
                restore()
            check('subprocess was invoked at all', len(spy.calls) == 1,
                  f'{len(spy.calls)} call(s)')
            cmd, kwargs = spy.calls[0]
            profile = profile_path(cmd)
            check('every invocation carries -env:UserInstallation=file://',
                  profile is not None)
            check('a timeout is always passed to soffice',
                  kwargs.get('timeout') == pdf_bridge.SOFFICE_TIMEOUT,
                  str(kwargs.get('timeout')))
            check('private profile dir was removed after the call',
                  profile is not None and not profile.exists(),
                  str(profile))
            check('the profile temp dir was removed too',
                  profile is not None and not profile.parent.exists(),
                  str(profile.parent))

        print('\n[3] timeout -> clean error, profile still cleaned up')
        spy = SpySubprocess()
        spy.raise_timeout = True
        restore = patch_subprocess(spy)
        try:
            expect_bridge_error(
                'timeout raises PdfBridgeError',
                lambda: pdf_to_pptx(SAMPLE, work, timeout=3),
                'timed out', status=504,
            )
        finally:
            restore()
        check('timed-out call still cleaned its profile dir',
              bool(spy.calls) and not profile_path(spy.calls[0][0]).exists(),
              str(profile_path(spy.calls[0][0]) if spy.calls else 'no call'))

        print('\n[4] non-zero exit -> error carries the stderr tail')
        spy = SpySubprocess()
        spy.returncode = 1
        spy.stderr = b'Error: no export filter'
        restore = patch_subprocess(spy)
        try:
            expect_bridge_error(
                'failed conversion raises PdfBridgeError with stderr',
                lambda: pptx_to_pdf(SAMPLE, work),
                'no export filter', status=500,
            )
        finally:
            restore()

        print('\n[5] missing binary -> clear error, not FileNotFoundError')
        original_soffice = pdf_bridge.soffice_path
        pdf_bridge.soffice_path = lambda: None
        try:
            expect_bridge_error(
                'missing soffice reports it',
                lambda: pdf_to_pptx(SAMPLE, work),
                'not installed', status=503,
            )
        finally:
            pdf_bridge.soffice_path = original_soffice

        print('\n[6] missing input file -> clear error')
        expect_bridge_error(
            'missing PDF reported',
            lambda: pdf_to_pptx(work / 'does-not-exist.pdf', work),
            'not found', status=500,
        )

        print('\n[7] concurrency: parallel conversions do not collide')
        if not soffice or not SAMPLE.exists():
            print('  (skipping: soffice or sample deck missing)')
        else:
            base = pptx_to_pdf(SAMPLE, work)
            srcs = []
            for tag in ('a', 'b', 'c'):
                src = work / f'deck-{tag}.pdf'
                shutil.copy(base, src)
                srcs.append(src)

            spy = SpySubprocess()
            restore = patch_subprocess(spy)
            results, errors = {}, []

            def worker(src):
                try:
                    results[src.name] = pdf_to_pptx(src, work / 'concurrent-out')
                except Exception as exc:  # noqa: BLE001 — reported as a failure
                    errors.append(f'{src.name}: {exc}')

            threads = [threading.Thread(target=worker, args=(src,)) for src in srcs]
            try:
                for th in threads:
                    th.start()
                for th in threads:
                    th.join()
            finally:
                restore()

            check('all concurrent conversions succeeded', not errors,
                  '; '.join(errors) or f'{len(results)} results')
            check('each conversion produced its own output',
                  len(set(results.values())) == len(srcs),
                  ', '.join(sorted(p.name for p in results.values())))
            check('outputs are all on disk',
                  all(p.exists() for p in results.values()))
            profiles = [profile_path(cmd) for cmd, _ in spy.calls]
            check('every concurrent call used a private profile',
                  all(p is not None for p in profiles), str(profiles))
            check('concurrent profiles never collided',
                  len({str(p) for p in profiles}) == len(spy.calls),
                  f'{len(set(map(str, profiles)))} distinct of {len(spy.calls)}')
            check('all concurrent profile dirs cleaned up',
                  all(p is not None and not p.exists() for p in profiles))

        print(f'\n{CHECKS - len(FAILURES)}/{CHECKS} checks passed')
        if FAILURES:
            print('FAILED:')
            for name in FAILURES:
                print(f'  - {name}')
            return 1
        print('PDF bridge contract holds.')
        return 0
    finally:
        shutil.rmtree(work, ignore_errors=True)


def test_pdf_bridge_contract() -> None:
    """`pytest tests/` runs the same checks as the script."""
    assert main() == 0, 'pdf-bridge checks failed'


if __name__ == '__main__':
    sys.exit(main())
