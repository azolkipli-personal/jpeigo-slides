#!/usr/bin/env python3
"""Why wasn't one paragraph's runs coalesced? Shows the extractor's own decision."""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path.home() / 'Projects/jpeigo-slide/backend'))

from pptx import Presentation  # noqa: E402
from app.core.extractor import _coalesce_groups, _rpr_signature  # noqa: E402

SRC = sys.argv[1] if len(sys.argv) > 1 else 'test_real.pptx'
prs = Presentation(SRC)


def walk(shapes, slide_idx, prefix=''):
    for idx, shape in enumerate(shapes):
        path = f'{prefix}{idx}'
        if shape.shape_type == 6:  # group
            walk(shape.shapes, slide_idx, prefix=f'{path}.')
            continue
        tf = getattr(shape, 'text_frame', None)
        if tf is None:
            continue
        for para_idx, para in enumerate(tf.paragraphs):
            runs = list(para.runs)
            if len(runs) < 2:
                continue
            groups = _coalesce_groups(runs)
            tag = 'MERGED' if len(groups) < len(runs) else 'LEFT PER-RUN'
            print(f'slide {slide_idx} shape {path} para {para_idx}: '
                  f'{len(runs)} runs -> {len(groups)} groups {groups}  [{tag}]')
            for i, r in enumerate(runs):
                sig = _rpr_signature(r) or '<none>'
                sig = sig.replace(' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"', '')
                print(f'   [{i}] {r.text!r:34} sig={sig[:110]}')


walk(prs.slides[int(sys.argv[2])].shapes, int(sys.argv[2]))
